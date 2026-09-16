# -*- coding: utf-8 -*-
"""把 public-apis 的全部 API 条目按用途分类导出为 Excel / CSV / JSON / Markdown。

数据来源：https://github.com/public-apis/public-apis  (MIT License)

用法：  python3 scripts/export_public_apis.py
输出：  data/public-apis-全部API清单.xlsx / .csv / 按用途分类.md / 按用途分类.json
"""
import csv
import json
import os
import re
from collections import OrderedDict

_HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.environ.get("PAP_BASE") or os.path.dirname(_HERE)
DATA = os.path.join(BASE, "data")


def _first(*cands):
    for c in cands:
        if c and os.path.exists(c):
            return c
    return [c for c in cands if c][0]


README = _first(os.environ.get("PAP_README"),
                os.path.join(DATA, "README-public-apis.md"),
                "D:/code/public-apis/README.md")
SRC_JSON = _first(os.environ.get("PAP_SRC_JSON"),
                  os.path.join(DATA, "public_apis.json"),
                  "D:/code/.workbuddy/tmp/public_apis.json")
OUT_DIR = os.environ.get("PAP_DATA") or DATA

COLUMNS = ["序号", "用途分类", "名称", "功能说明", "鉴权方式", "HTTPS", "CORS", "链接"]

# 鉴权方式的中文解释
AUTH_ZH = {
    "No": "无需鉴权（直接可用）",
    "apiKey": "需要 API Key",
    "OAuth": "需要 OAuth 授权",
    "X-Mashape-Key": "需要 X-Mashape-Key",
    "User-Agent": "需要自定义 User-Agent",
}


def load_records():
    recs = json.load(open(SRC_JSON, encoding="utf-8"))
    # 补录 APILayer 赞助区块的 8 条（README 第 40-60 行附近的赞助表格）
    lines = open(README, encoding="utf-8", errors="replace").read().split("\n")
    header = None
    for idx, line in enumerate(lines, 1):
        if not (11 <= idx <= 60):
            continue
        s = line.strip()
        if "|" not in s:
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if cells[0].lower() == "api" and "description" in [c.lower() for c in cells]:
            header = cells
            continue
        if header is None or set("".join(cells)) <= set(":- "):
            continue
        m = re.match(r"^\|\s*\[([^\]]+)\]\(([^)]+)\)", s)
        if not m:
            continue
        row = {"name": m.group(1).strip(), "url": m.group(2).strip(),
               "category": "APILayer Suite（赞助）"}
        for i, col in enumerate(header):
            if i == 0 or i >= len(cells):
                continue
            key = col.lower().strip()
            key = {"auth": "auth", "https": "https", "cors": "cors",
                   "description": "description"}.get(key, key)
            row[key] = re.sub(r"\s+", " ", cells[i].replace("`", "")).strip()
        for k in ("description", "auth", "https", "cors"):
            row.setdefault(k, "")
        recs.append(row)
    return recs


def build_stats(recs, order):
    stats = OrderedDict()
    for cat in order:
        rows = [r for r in recs if r["category"] == cat]
        stats[cat] = {
            "total": len(rows),
            "no_auth": sum(1 for r in rows if r["auth"].lower() == "no"),
            "api_key": sum(1 for r in rows if r["auth"].lower() == "apikey"),
            "oauth": sum(1 for r in rows if r["auth"].lower() == "oauth"),
            "https": sum(1 for r in rows if r["https"].lower() == "yes"),
            "cors_yes": sum(1 for r in rows if r["cors"].lower() == "yes"),
        }
    return stats


def write_csv(recs, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for i, r in enumerate(recs, 1):
            w.writerow([i, r["category"], r["name"], r["description"],
                        r["auth"], r["https"], r["cors"], r["url"]])


def write_grouped_json(recs, order, path):
    data = OrderedDict()
    for cat in order:
        data[cat] = [
            {"name": r["name"], "description": r["description"], "auth": r["auth"],
             "https": r["https"], "cors": r["cors"], "url": r["url"]}
            for r in recs if r["category"] == cat
        ]
    json.dump(data, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=2)


def write_markdown(recs, stats, order, path):
    lines = []
    lines.append("# Public APIs 全量清单（按用途分类）\n")
    lines.append("> 来源：`public-apis/public-apis` 仓库 README\n")
    lines.append("> 共 **%d** 个用途分类、**%d** 个 API 条目\n" % (len(order), len(recs)))
    lines.append("")
    lines.append("## 用途总览\n")
    lines.append("| 用途分类 | 条目数 | 免鉴权 | 需 API Key | 支持 HTTPS |")
    lines.append("|---|---:|---:|---:|---:|")
    for cat in order:
        s = stats[cat]
        anchor = cat.lower().replace(" & ", "--").replace(" ", "-")
        lines.append("| [%s](#%s) | %d | %d | %d | %d |"
                     % (cat, anchor, s["total"], s["no_auth"], s["api_key"], s["https"]))
    lines.append("")
    lines.append("## 完整清单\n")
    for cat in order:
        rows = [r for r in recs if r["category"] == cat]
        lines.append("### %s\n" % cat)
        lines.append("| 名称 | 功能说明 | 鉴权 | HTTPS | CORS | 链接 |")
        lines.append("|---|---|---|---|---|---|")
        for r in rows:
            desc = r["description"].replace("|", "\\|")
            lines.append("| **%s** | %s | %s | %s | %s | [打开](%s) |"
                         % (r["name"], desc, r["auth"] or "-", r["https"] or "-",
                            r["cors"] or "-", r["url"]))
        lines.append("")
    open(path, "w", encoding="utf-8").write("\n".join(lines))


def write_xlsx(recs, stats, order, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
    HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
    ALT_FILL = PatternFill("solid", fgColor="EAF2FA")
    FREE_FILL = PatternFill("solid", fgColor="E3F5E6")
    THIN = Side(style="thin", color="C9D6E3")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    wb = Workbook()

    # ---------- Sheet 1: 全部 API ----------
    ws = wb.active
    ws.title = "全部API"
    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, r in enumerate(recs, 1):
        ws.append([i, r["category"], r["name"], r["description"],
                   r["auth"], r["https"], r["cors"], r["url"]])
        row = ws.max_row
        free = r["auth"].lower() == "no"
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center",
                                       wrap_text=(c == 4),
                                       horizontal="center" if c in (1, 5, 6, 7) else "left")
            if free:
                cell.fill = FREE_FILL
            elif i % 2 == 0:
                cell.fill = ALT_FILL
        link = ws.cell(row=row, column=8)
        link.hyperlink = r["url"]
        link.font = Font(color="0563C1", underline="single", size=10)
        ws.cell(row=row, column=3).font = Font(bold=True, size=10)
    widths = [6, 24, 26, 62, 13, 8, 10, 42]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:H%d" % ws.max_row

    # ---------- Sheet 2: 用途总览 ----------
    ws2 = wb.create_sheet("用途总览")
    ws2.append(["用途分类", "条目数", "免鉴权（开箱可用）", "需 API Key", "需 OAuth",
                "支持 HTTPS", "支持 CORS", "免鉴权占比"])
    for c in range(1, 9):
        cell = ws2.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cat in order:
        s = stats[cat]
        ws2.append([cat, s["total"], s["no_auth"], s["api_key"], s["oauth"],
                    s["https"], s["cors_yes"],
                    round(s["no_auth"] / s["total"] * 100, 1) if s["total"] else 0])
        row = ws2.max_row
        for c in range(1, 9):
            cell = ws2.cell(row=row, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if c > 1 else "left",
                                       vertical="center")
        ws2.cell(row=row, column=8).number_format = "0.0\"%\""
    tot = len(recs)
    ws2.append(["合计", tot, sum(s["no_auth"] for s in stats.values()),
                sum(s["api_key"] for s in stats.values()),
                sum(s["oauth"] for s in stats.values()),
                sum(s["https"] for s in stats.values()),
                sum(s["cors_yes"] for s in stats.values()),
                round(sum(s["no_auth"] for s in stats.values()) / tot * 100, 1)])
    for c in range(1, 9):
        cell = ws2.cell(row=ws2.max_row, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for i, w in enumerate([26, 10, 20, 12, 12, 12, 12, 12], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    # ---------- Sheet 3: 免鉴权精选 ----------
    ws3 = wb.create_sheet("免鉴权精选")
    ws3.append(["序号", "用途分类", "名称", "功能说明", "HTTPS", "CORS", "链接"])
    for c in range(1, 8):
        cell = ws3.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    free_rows = [r for r in recs
                 if r["auth"].lower() == "no" and r["https"].lower() == "yes"]
    for i, r in enumerate(free_rows, 1):
        ws3.append([i, r["category"], r["name"], r["description"],
                    r["https"], r["cors"], r["url"]])
        row = ws3.max_row
        for c in range(1, 8):
            cell = ws3.cell(row=row, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 4))
        link = ws3.cell(row=row, column=7)
        link.hyperlink = r["url"]
        link.font = Font(color="0563C1", underline="single", size=10)
    for i, w in enumerate([6, 24, 26, 62, 8, 10, 42], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = "A1:G%d" % ws3.max_row

    # ---------- Sheet 4: 说明 ----------
    ws4 = wb.create_sheet("使用说明")
    notes = [
        ["Public APIs 全量清单", ""],
        ["数据来源", "https://github.com/public-apis/public-apis （README，master 分支）"],
        ["整理时间", "2026-09-16"],
        ["用途分类数", len(order)],
        ["API 总条目", len(recs)],
        ["其中免鉴权可开箱使用", sum(1 for r in recs if r["auth"].lower() == "no")],
        ["其中支持 HTTPS", sum(1 for r in recs if r["https"].lower() == "yes")],
        ["", ""],
        ["字段说明", ""],
        ["鉴权方式", "No=无需鉴权 / apiKey=需 API Key / OAuth=需 OAuth 授权"],
        ["HTTPS", "Yes=支持 HTTPS（生产环境建议优先选）"],
        ["CORS", "Yes=允许浏览器跨域直连 / Unknown=官方未说明 / No=不支持跨域"],
        ["", ""],
        ["使用建议", ""],
        ["① 优先「免鉴权精选」表", "无需注册即可直接调用，适合原型和 Demo"],
        ["② 生产环境优先 HTTPS=Yes", "避免混合内容被浏览器拦截"],
        ["③ 纯前端调用需 CORS=Yes", "后端调用则 CORS 无关紧要"],
        ["④ 调用前请核对官方条款", "免费额度、速率限制、商用授权请以官网为准"],
    ]
    for r in notes:
        ws4.append(r)
    ws4.column_dimensions["A"].width = 30
    ws4.column_dimensions["B"].width = 76
    ws4["A1"].font = Font(bold=True, size=14, color="1F4E79")
    for row in ws4.iter_rows(min_row=9, max_row=9):
        for c in row:
            if c.value:
                c.font = Font(bold=True, color="1F4E79")
    for row in ws4.iter_rows(min_row=14, max_row=14):
        for c in row:
            if c.value:
                c.font = Font(bold=True, color="1F4E79")

    wb.save(path)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    recs = load_records()
    order = []
    for r in recs:
        if r["category"] not in order:
            order.append(r["category"])
    stats = build_stats(recs, order)

    write_csv(recs, os.path.join(OUT_DIR, "public-apis-全部API清单.csv"))
    write_grouped_json(recs, order, os.path.join(OUT_DIR, "public-apis-按用途分类.json"))
    write_markdown(recs, stats, order, os.path.join(OUT_DIR, "public-apis-按用途分类.md"))
    write_xlsx(recs, stats, order, os.path.join(OUT_DIR, "public-apis-全部API清单.xlsx"))

    print("records=%d categories=%d" % (len(recs), len(order)))
    print("free_no_auth=%d" % sum(1 for r in recs if r["auth"].lower() == "no"))
    print("https_yes=%d" % sum(1 for r in recs if r["https"].lower() == "yes"))
    for f in sorted(os.listdir(OUT_DIR)):
        print("  %-42s %.1f KB" % (f, os.path.getsize(os.path.join(OUT_DIR, f)) / 1024))


if __name__ == "__main__":
    main()
